"""XLSX-Export: schreibt ein ExportPayload als echte .xlsx-Datei.

Formel-Maskierung ("=" -> "'=") erfolgt ausschließlich hier, unmittelbar vor
dem Zellschreiben - Vorschau und PDF-Export zeigen weiterhin den Rohwert
(siehe DESIGN.md, 6.2).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.export.units import px_to_excel_width, px_to_points
from core.models.export_payload import ExportPayload

ALTERNATE_ROW_FILL = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
HEADER_ROW = 1
FIRST_DATA_ROW = 2


def _mask_formula_like(value: str) -> str:
    """Verhindert, dass Excel einen mit '=' beginnenden Textwert als Formel interpretiert."""
    return "'" + value if value.startswith("=") else value


_INVALID_SHEET_CHARS = "[]:*?/\\"
_MAX_SHEET_NAME = 31


def unique_sheet_name(name: str, used: set[str]) -> str:
    """Excel-gültiger, eindeutiger Reitername (max. 31 Zeichen, ohne unzulässige Zeichen)."""
    cleaned = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in name).strip("' ") or "Tabelle"
    candidate = cleaned[:_MAX_SHEET_NAME]
    counter = 2
    while candidate.lower() in used:
        suffix = f"_{counter}"
        candidate = cleaned[:_MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


class XlsxExporter:
    def export(self, payload: ExportPayload, target_path: Path) -> None:
        self.export_many([("Export", payload)], target_path)

    def export_many(self, sheets: list[tuple[str, ExportPayload]], target_path: Path) -> None:
        """Schreibt eine Arbeitsmappe mit einem Reiter je (Name, Payload)."""
        workbook = Workbook()
        for position, (name, payload) in enumerate(sheets):
            self.add_sheet(workbook, name, payload, reuse_active=position == 0)

        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target_path)

    def add_sheet(self, workbook: Workbook, name: str, payload: ExportPayload, reuse_active: bool = False) -> None:
        """Fügt einen Reiter hinzu (reuse_active: das leere Standardblatt der neuen Mappe nutzen)."""
        worksheet = workbook.active if reuse_active else workbook.create_sheet()
        used = {ws.title.lower() for ws in workbook.worksheets if ws is not worksheet}
        worksheet.title = unique_sheet_name(name, used)
        self._write_print_header(worksheet, payload)
        self._write_column_headers(worksheet, payload)
        self._write_rows(worksheet, payload)
        self._apply_column_widths(worksheet, payload)
        worksheet.page_setup.orientation = "landscape" if payload.landscape else "portrait"

    @staticmethod
    def _write_print_header(worksheet: Worksheet, payload: ExportPayload) -> None:
        worksheet.oddHeader.left.text = " | ".join(payload.header.left) or None
        worksheet.oddHeader.center.text = " | ".join(payload.header.center) or None
        worksheet.oddHeader.right.text = " | ".join(payload.header.right) or None
        worksheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    @staticmethod
    def _write_column_headers(worksheet: Worksheet, payload: ExportPayload) -> None:
        font = Font(size=payload.header_font_size, bold=True)
        worksheet.row_dimensions[HEADER_ROW].height = px_to_points(payload.header_row_height_px)
        for col_index, column in enumerate(payload.columns, start=1):
            cell = worksheet.cell(row=HEADER_ROW, column=col_index, value=column.name)
            cell.font = font

    def _write_rows(self, worksheet: Worksheet, payload: ExportPayload) -> None:
        font = Font(size=payload.font_size)
        alignment = Alignment(wrap_text=payload.word_wrap)
        data_row_counter = 0
        for row_offset, row in enumerate(payload.rows, start=FIRST_DATA_ROW):
            if row.is_separator:
                worksheet.row_dimensions[row_offset].height = px_to_points(payload.separator_height_px)
                continue

            worksheet.row_dimensions[row_offset].height = px_to_points(payload.row_height_px)
            use_alternate = payload.alternate_row_colors and data_row_counter % 2 == 1
            for col_index, value in enumerate(row.values, start=1):
                cell = worksheet.cell(row=row_offset, column=col_index, value=_mask_formula_like(value))
                cell.font = font
                cell.alignment = alignment
                fill = self._resolve_cell_fill(row, col_index - 1, use_alternate)
                if fill is not None:
                    cell.fill = fill
            data_row_counter += 1

    @staticmethod
    def _resolve_cell_fill(row, col_index: int, use_alternate: bool):
        # Zeilenfarbe (Reiter 3, Phase 2) hat höchste Priorität
        if row.row_color:
            hex_val = row.row_color.lstrip("#")
            return PatternFill(start_color=hex_val, end_color=hex_val, fill_type="solid")
        # Zellenfarbe (Reiter 3, Phase 1)
        if row.cell_colors and col_index < len(row.cell_colors):
            cell_color = row.cell_colors[col_index]
            if cell_color:
                hex_val = cell_color.lstrip("#")
                return PatternFill(start_color=hex_val, end_color=hex_val, fill_type="solid")
        # Alternierende Zeilenfarbe
        if use_alternate:
            return ALTERNATE_ROW_FILL
        return None

    @staticmethod
    def _apply_column_widths(worksheet: Worksheet, payload: ExportPayload) -> None:
        for col_index, column in enumerate(payload.columns, start=1):
            letter = get_column_letter(col_index)
            worksheet.column_dimensions[letter].width = px_to_excel_width(column.width_px)
