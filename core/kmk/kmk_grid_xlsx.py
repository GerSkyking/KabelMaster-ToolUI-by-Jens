"""Excel-Rasteransicht der KMK-Beschrifter: gleiche Anordnung wie im PDF.

Die Platzierung (Hoch-/Querformat nach Kapazität, zeilenweise Füllung, Überschrift-Modi
"first_field"/"own_row"/"hidden", nahtlose Gruppen, Seitenumbrüche) spiegelt
KmkPdfGenerator. Jede Zelle ist ein Beschrifter; Zeilenhöhen/Spaltenbreiten entsprechen
dem KMK-Typ, sodass die Excel-Seite dem PDF-Blatt entspricht.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from core.export.xlsx_exporter import XlsxExporter
from core.models.export_payload import ExportPayload
from core.models.kmk_models import MM_TO_PT, KmkConfig, KmkGroup

MARGIN_MM = 10.0
A4_MM = (210.0, 297.0)
_EXCEL_MM_PER_CHAR = 1.85  # ca. Breite eines Zeichens der Standardschrift in mm
_HEADER_FILL = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")


@dataclass
class GridCell:
    page: int
    row: int
    col: int
    lines: list[str]
    is_header: bool = False
    span: int = 1  # >1: Überschrift über die ganze Zeile (own_row)


def compute_grid(config: KmkConfig) -> tuple[bool, int, int]:
    """(Querformat, Spalten, Zeilen) pro Seite - identisch zur Wahl im PDF."""
    label_w_mm, label_h_mm = config.kmk_type.width_mm, config.kmk_type.height_mm

    def fit(page_w: float, page_h: float) -> tuple[int, int]:
        return (
            max(1, int((page_w - 2 * MARGIN_MM) / label_w_mm)),
            max(1, int((page_h - 2 * MARGIN_MM) / label_h_mm)),
        )

    p_cols, p_rows = fit(*A4_MM)
    l_cols, l_rows = fit(A4_MM[1], A4_MM[0])
    if l_cols * l_rows > p_cols * p_rows:
        return True, l_cols, l_rows
    return False, p_cols, p_rows


def layout_cells(groups: list[KmkGroup], config: KmkConfig) -> tuple[list[GridCell], bool, int, int]:
    landscape, cols, rows = compute_grid(config)
    cells: list[GridCell] = []
    page = col = row = 0

    def check_new_page() -> None:
        nonlocal page, row
        if row >= rows:
            page += 1
            row = 0

    def complete_row() -> None:
        nonlocal col, row
        if col > 0:
            col = 0
            row += 1
            check_new_page()

    def place(lines: list[str], is_header: bool = False) -> None:
        nonlocal col, row
        cells.append(GridCell(page, row, col, lines, is_header))
        col += 1
        if col >= cols:
            col = 0
            row += 1
            check_new_page()

    for group in groups:
        if config.header_mode == "own_row" or not config.seamless:
            complete_row()
            check_new_page()
        title = f"<< {group.title} ({len(group.labels)} Kabel) >>"
        if config.header_mode == "own_row":
            cells.append(GridCell(page, row, 0, [title], True, span=cols))
            col = 0
            row += 1
            check_new_page()
        elif config.header_mode == "first_field":
            place([title], is_header=True)
        for label in group.labels:
            place([ln for ln in label.lines if ln.strip()])
    return cells, landscape, cols, rows


def _text(value: str) -> str:
    return value[1:] if value.startswith("'=") else value


class KmkGridXlsx:
    """Schreibt Raster-Blatt ("Raster") und optional die Kontrollliste ("Liste") in eine Arbeitsmappe."""

    def write(
        self,
        groups: list[KmkGroup],
        config: KmkConfig,
        target_path: Path,
        list_payload: ExportPayload | None = None,
    ) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Raster"
        self._fill_grid(sheet, groups, config)
        if list_payload is not None:
            XlsxExporter().add_sheet(workbook, "Liste", list_payload)
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target_path)

    @staticmethod
    def _fill_grid(sheet, groups: list[KmkGroup], config: KmkConfig) -> None:
        cells, landscape, cols, rows = layout_cells(groups, config)
        pages = (max((c.page for c in cells), default=0) + 1) if cells else 1

        width_chars = config.kmk_type.width_mm / _EXCEL_MM_PER_CHAR + 0.71
        for col in range(cols):
            sheet.column_dimensions[get_column_letter(col + 1)].width = width_chars
        for row in range(pages * rows):
            sheet.row_dimensions[row + 1].height = config.kmk_type.height_mm * MM_TO_PT

        color = config.frame_color.lstrip("#") or "000000"
        side = Side(style="thin", color=color)
        frame = Border(left=side, right=side, top=side, bottom=side)
        black_side = Side(style="thin", color="000000")
        header_frame = Border(left=black_side, right=black_side, top=black_side, bottom=black_side)
        font = Font(size=config.font_size_pt)
        bold = Font(size=config.font_size_pt, bold=True)
        centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for cell in cells:
            excel_row = cell.page * rows + cell.row + 1
            target = sheet.cell(
                row=excel_row, column=cell.col + 1, value="\n".join(_text(line) for line in cell.lines)
            )
            target.alignment = centered
            target.font = bold if cell.is_header else font
            if cell.is_header:
                target.fill = _HEADER_FILL
                target.border = header_frame
            elif config.frame_enabled:
                target.border = frame
            if cell.span > 1:
                sheet.merge_cells(start_row=excel_row, start_column=1, end_row=excel_row, end_column=cell.span)

        for page in range(1, pages):
            sheet.row_breaks.append(Break(id=page * rows))
        sheet.page_setup.orientation = "landscape" if landscape else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        margin_inch = MARGIN_MM / 25.4
        sheet.page_margins.left = sheet.page_margins.right = margin_inch
        sheet.page_margins.top = sheet.page_margins.bottom = margin_inch
